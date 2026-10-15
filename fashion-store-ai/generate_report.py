# -*- coding: utf-8 -*-
import sys
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_report():
    doc = docx.Document('e:/doan/fashion-store-ai/BaocaoM1-D2-Nguyễn Đức Chí-0002767-67PM.docx')

    # Update Table 1: Thông tin sinh viên & đề tài
    t1 = doc.tables[1]
    t1.rows[0].cells[1].text = "Phạm Ngọc Đức"
    t1.rows[0].cells[3].text = "0288767"
    t1.rows[1].cells[1].text = "67PM3"
    t1.rows[1].cells[3].text = "K67"
    t1.rows[2].cells[1].text = "Công nghệ thông tin"
    t1.rows[2].cells[3].text = "Cử nhân"
    
    # Row 3: Tên đề tài
    topic = "Hệ thống cửa hàng thời trang trực tuyến tích hợp chatbot AI hỗ trợ khách hàng và quản lý bán hàng"
    for c in t1.rows[3].cells[1:]:
        c.text = topic
        
    t1.rows[4].cells[1].text = "Công nghệ phần mềm"
    t1.rows[4].cells[3].text = "Công nghệ phần mềm"
    t1.rows[5].cells[1].text = "Phan Hữu Trung"
    t1.rows[5].cells[3].text = "ducp9987789@gmail.com"
    t1.rows[6].cells[1].text = "03/08/2026 – 29/11/2026"
    t1.rows[6].cells[3].text = "Đợt 1"

    # Set font Times New Roman 13pt for Table 1
    for row in t1.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(13)

    # Update Paragraphs in Section 1, 2, 3
    # Let's see all paragraphs and update them appropriately
    
    # 1. Mục tiêu và yêu cầu của ĐATN
    p_map = {
        7: "Xây dựng một nền tảng thương mại điện tử chuyên ngành thời trang trực tuyến hoàn chỉnh, hiện đại, tối ưu trải nghiệm mua sắm cho khách hàng và cung cấp bộ công cụ quản lý bán hàng mạnh mẽ cho người quản trị.",
        8: "Triển khai kiến trúc Client - Server - AI Service hiện đại: Frontend SPA sử dụng Vue 3 (Pinia, Vue Router), Backend RESTful API sử dụng Python FastAPI kết hợp cơ sở dữ liệu MongoDB, đảm bảo khả năng mở rộng và tốc độ phản hồi nhanh.",
        9: "Tích hợp Trợ lý ảo Chatbot AI Stylist thông minh hoạt động 24/7 ứng dụng kiến trúc RAG (Retrieval-Augmented Generation) kết hợp Vector Database (ChromaDB) và mô hình ngôn ngữ lớn (Gemini / OpenAI LLM) nhằm tư vấn trang phục, gợi ý size chuẩn xác và giải đáp thông tin sản phẩm tức thì.",
        10: "< Dự kiến > Cơ chế AI Visual Search & Smart Recommendation: Nghiên cứu và mở rộng tính năng gợi ý sản phẩm nâng cao dựa trên sở thích, lịch sử tương tác và hình ảnh mẫu của người dùng. Hệ thống kết hợp phân tích ngữ cảnh người dùng để cá nhân hóa danh mục gợi ý, tạo trải nghiệm mua sắm thông minh và đột phá.",
        
        12: "Hệ thống Mua sắm (E-commerce Client): Cung cấp đầy đủ các chức năng xem sản phẩm, bộ lọc danh mục độc lập (khoảng giá, màu sắc, kích cỡ, danh mục), tìm kiếm thời gian thực, xem chi tiết sản phẩm, quản lý giỏ hàng, đặt hàng và theo dõi trạng thái đơn hàng.",
        13: "Quản lý Tài khoản & Phân quyền: Cung cấp chức năng Đăng ký, Đăng nhập (JWT), xác thực thông tin tài khoản, phân quyền truy cập giữa Khách hàng (Customer) và Quản trị viên (Admin).",
        14: "Trợ lý ảo Chatbot AI Stylist: Tích hợp khung chat tương tác trực tiếp trên giao diện web, có khả năng trả lời tự nhiên các câu hỏi về sản phẩm, tư vấn phối đồ outfit theo nhu cầu/sự kiện, tư vấn kích thước theo chiều cao và cân nặng, tra cứu mã giảm giá và chính sách cửa hàng dựa trên tri thức RAG.",
        15: "Hệ thống Quản trị Bán hàng (Admin Dashboard): Xây dựng bảng điều khiển trực quan cho quản trị viên gồm thống kê doanh thu, đơn hàng; Quản lý sản phẩm (CRUD, phân loại màu/size, giá khuyến mãi); Quản lý danh mục thời trang; Quản lý đơn hàng; Quản lý chương trình khuyến mãi (Voucher); Quản lý thông tin cửa hàng.",
        16: "Quản lý Dữ liệu Tri thức Chatbot: Cung cấp giao diện quản trị cho phép cập nhật dữ liệu sản phẩm, chính sách đổi trả, câu hỏi thường gặp (FAQ) và đồng bộ hóa tự động vào Vector Database ChromaDB.",
        17: "Hệ thống Quản lý Khuyến mãi & Voucher: Quản lý các mã giảm giá theo phần trăm hoặc số tiền cố định, áp dụng tức thì tại bước thanh toán.",
        18: "Quản lý người dùng: Theo dõi danh sách tài khoản, trạng thái hoạt động và lịch sử mua sắm của khách hàng.",
        19: "Quản lý Sản phẩm & Danh mục: Thêm, sửa, xóa, cấu hình thuộc tính sản phẩm (màu sắc, bảng size, hình ảnh, giá gốc, giá sale) và tổ chức cây danh mục linh hoạt.",
        20: "Quản lý Đơn hàng & Vận chuyển: Theo dõi trạng thái đơn hàng (Chờ xác nhận, Đang giao, Đã hoàn thành, Đã hủy) và cập nhật thông tin vận chuyển.",

        22: "Kiến trúc: Phát triển theo mô hình phân tầng Client - Server - AI Service với Frontend SPA (Vue 3, Vite) và Backend RESTful API (Python FastAPI, Motor Async MongoDB Driver), đảm bảo sự độc lập và khả năng mở rộng cao.",
        23: "Hiệu năng & Trải nghiệm: Tốc độ phản hồi API < 200ms; giao diện Responsive tối ưu hiển thị mượt mà trên Desktop, Tablet và Mobile; bộ lọc danh mục dạng sticky cuộn độc lập giúp tối ưu trải nghiệm duyệt sản phẩm.",
        24: "An toàn & Bảo mật: Xác thực người dùng bằng JSON Web Token (JWT); mã hóa mật khẩu an toàn bằng thuật toán Bcrypt; phân quyền truy cập nghiêm ngặt giữa khách hàng và admin; bảo vệ các endpoint API quan trọng.",
        25: "Công nghệ & Chất lượng mã nguồn: Sử dụng ngôn ngữ Python 3.11+, JavaScript (ES6+), framework Vue 3 và FastAPI; quản lý mã nguồn theo cấu trúc module hóa, tuân thủ tiêu chuẩn PEP8 và ESLint, dễ bảo trì và mở rộng.",
        26: "Đóng gói & Triển khai: Đóng gói toàn bộ hệ thống bằng Docker và Docker Compose (Frontend, Backend, MongoDB, ChromaDB), đảm bảo tính nhất quán và dễ dàng triển khai trên mọi môi trường.",

        # 2. Phạm vi, nội dung công việc và sản phẩm cần nộp
        30: "Hệ thống Web E-commerce: Phát triển nền tảng thương mại điện tử trực tuyến chạy trên trình duyệt web, tập trung vào trải nghiệm mua sắm thời trang mượt mà, tiện lợi và giao diện hiện đại.",
        31: "Hệ thống Trợ lý ảo AI: Xây dựng dịch vụ Chatbot AI Stylist thông minh sử dụng RAG kết hợp ChromaDB và LLM để tư vấn và hỗ trợ khách hàng 24/7.",
        32: "Hệ thống Quản trị Admin: Xây dựng phân hệ quản lý tập trung dành cho ban quản trị để điều phối sản phẩm, đơn hàng, danh mục, khuyến mãi và tri thức chatbot.",
        34: "Frontend (Vue 3, Pinia, Vite): Xây dựng các trang chức năng (Trang chủ, Danh mục, Chi tiết sản phẩm, Giỏ hàng, Đặt hàng, Tra cứu đơn hàng, Admin Dashboard) và Chat Widget tương tác AI.",
        35: "Backend (Python, FastAPI): Thiết kế hệ thống RESTful API xử lý xác thực người dùng, nghiệp vụ bán hàng, quản lý đơn hàng và tích hợp dịch vụ AI Core.",
        36: "Database (MongoDB & ChromaDB): Thiết kế cơ sở dữ liệu NoSQL MongoDB lưu trữ dữ liệu thực thể (users, products, categories, orders, carts, promotions, chat_logs) và Vector Database ChromaDB lưu trữ embeddings tri thức phục vụ RAG.",
        37: "Admin Dashboard: Module quản trị hỗ trợ giao diện CRUD chuyên nghiệp, bảng thống kê số liệu và quản lý dữ liệu sản phẩm, đơn hàng và FAQ chatbot.",
        38: "< Dự kiến > AI Recommendation Engine: Nghiên cứu tích hợp cơ chế gợi ý sản phẩm thông minh dựa trên độ tương đồng vector và thói quen xem sản phẩm của người dùng.",

        40: "Nghiên cứu yêu cầu nghiệp vụ thương mại điện tử thời trang, khảo sát công nghệ AI/RAG, thiết kế sơ đồ Use Case, Activity Diagram và kiến trúc hệ thống Client - Server - AI Service.",
        41: "Phát triển khung giao diện Frontend Vue 3, hệ thống định tuyến Vue Router, quản lý trạng thái tập trung với Pinia và bộ lọc danh mục độc lập.",
        42: "Xây dựng Backend RESTful API với FastAPI, triển khai bảo mật JWT, mã hóa Bcrypt và thiết lập kết nối bất đồng bộ với MongoDB.",
        43: "Xây dựng dịch vụ RAG AI Stylist: pipeline chuyển đổi embedding, lưu trữ vector vào ChromaDB, xây dựng Prompt Engineering và kết nối LLM (Gemini/OpenAI).",
        44: "Phát triển phân hệ Admin Dashboard để quản trị toàn diện hệ thống bán hàng và cập nhật tri thức cho chatbot.",
        45: "Kiểm thử hệ thống, tối ưu hóa tốc độ tải trang, đóng gói ứng dụng với Docker Compose và hoàn thiện tài liệu báo cáo ĐATN.",

        47: "Mã nguồn: Toàn bộ mã nguồn Frontend (Vue 3), Backend (FastAPI), AI Service và cấu hình Docker Compose được tổ chức khoa học trên Git repository.",
        48: "Tài liệu: Báo cáo đồ án tốt nghiệp, sơ đồ kiến trúc hệ thống, biểu đồ UML, tài liệu đặc tả API và hướng dẫn cài đặt, vận hành hệ thống.",
        49: "Cơ sở dữ liệu: File export cấu trúc và bộ dữ liệu mẫu (seed data) cho MongoDB và Vector Collection trong ChromaDB.",
        50: "Bản demo vận hành: Hệ thống hoạt động hoàn chỉnh trên môi trường Docker, cho phép đăng ký, duyệt sản phẩm, lọc thông minh, đặt hàng, quản trị và chat với AI Stylist.",
        51: "Video demo: Video quay lại quy trình sử dụng các tính năng chính của khách hàng, tương tác tư vấn với AI Stylist và thao tác quản trị trên Admin Dashboard.",
        52: "Bảo vệ: Slide thuyết trình, Poster giới thiệu đề tài và triển khai thử nghiệm hệ thống lên môi trường Internet/Cloud.",

        # 3. Dữ liệu đầu vào, giả thiết, tiêu chuẩn...
        55: "Dữ liệu danh mục & sản phẩm: Dữ liệu mẫu phong phú về các mặt hàng thời trang (áo sơ mi, áo thun, quần âu, chân váy, phụ kiện...) kèm hình ảnh, giá bán, màu sắc, kích thước và tồn kho.",
        56: "Bộ tri thức Chatbot (RAG): Tập hợp tài liệu FAQ, chính sách giao hàng, quy định đổi trả 30 ngày, bảng quy chuẩn size theo số đo cơ thể để nạp vào ChromaDB.",
        57: "Dữ liệu người dùng & đơn hàng: Các tài khoản mẫu khách hàng và admin, kèm lịch sử đơn hàng với nhiều trạng thái để kiểm thử quy trình xử lý đơn hàng và bảng thống kê.",
        58: "Tài nguyên đa phương tiện: Hình ảnh sản phẩm chất lượng cao, icon giao diện dạng SVG và font chữ tối ưu để kiểm thử hiệu năng hiển thị và tốc độ tải trang.",

        60: "Phần cứng & Môi trường: Hệ thống được phát triển và kiểm thử trên máy tính cá nhân (CPU Intel Core i5/i7, 8GB-16GB RAM), đóng gói bằng Docker và vận hành ổn định trên mọi trình duyệt web hiện đại.",
        61: "Điều kiện phát triển: Đồ án thực hiện độc lập, sử dụng các công nghệ mã nguồn mở hiện đại (Vue 3, FastAPI, MongoDB, ChromaDB) và kết nối API mô hình ngôn ngữ lớn (Gemini API / OpenAI API).",
        62: "Khả năng mở rộng: Kiến trúc module hóa cho phép dễ dàng tích hợp cổng thanh toán trực tuyến (VNPAY, MoMo), mở rộng hệ thống bán hàng đa kênh hoặc nâng cấp thuật toán gợi ý AI.",

        64: "Lập trình: Áp dụng chuẩn JavaScript ES6+, Vue 3 Composition API, Python PEP8, Pydantic Type Annotations và chuẩn thiết kế RESTful API.",
        65: "Thiết kế: Tuân thủ nguyên lý thiết kế UI/UX hiện đại, giao diện sạch sẽ, thân thiện, tương thích đa thiết bị (Responsive Design).",
        66: "Bảo mật: Chuẩn mã hóa mật khẩu Bcrypt, xác thực JWT Bearer Token, kiểm soát phân quyền RBAC và chống lỗ hổng Injection.",
        67: "Quản lý phiên bản: Sử dụng Git để quản lý mã nguồn, chia nhánh tính năng và lưu trữ lịch sử phát triển dự án.",

        69: "Frontend: Sử dụng Vue 3, Vite, Pinia, Vue Router, Axios để xây dựng giao diện Single Page Application hiện đại và tương tác mượt mà.",
        70: "Backend: Xây dựng trên nền tảng Python 3.11+, framework FastAPI, Uvicorn, Motor async driver cho MongoDB, xử lý đa tác vụ với hiệu năng cao.",
        71: "Cơ sở dữ liệu: Sử dụng MongoDB để lưu trữ dữ liệu thực thể và ChromaDB để lưu trữ vector embeddings phục vụ truy xuất ngữ cảnh cho Chatbot AI.",
        72: "< Dự kiến > AI Core & RAG Pipeline: Tích hợp mô hình ngôn ngữ lớn (Google Gemini / OpenAI API) kết hợp ChromaDB để xử lý hội thoại thông minh và tư vấn thời trang.",
        74: "Sử dụng Visual Studio Code làm môi trường lập trình chính (IDE).",
        75: "Sử dụng Postman để thiết kế, kiểm tra và thử nghiệm các endpoint RESTful API.",
        76: "Sử dụng Docker và Docker Compose để container hóa và đồng bộ môi trường triển khai.",
        77: "Kiểm thử: Sử dụng Pytest và các công cụ DevTools trên trình duyệt để kiểm thử chức năng, tính đúng đắn của dữ liệu và hiệu năng tải trang."
    }

    for p_idx, text in p_map.items():
        if p_idx < len(doc.paragraphs):
            p = doc.paragraphs[p_idx]
            p.text = text
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(13)

    # Update Table 2: Kế hoạch/mốc tiến độ dự kiến (7 mốc từ 03/08/2026 đến 29/11/2026)
    t2 = doc.tables[2]
    milestones = [
        ("1", "Nghiên cứu đề tài hệ thống thương mại điện tử thời trang tích hợp Chatbot AI, khảo sát công nghệ (Vue 3, FastAPI, RAG, ChromaDB), phân tích yêu cầu hệ thống, xây dựng sơ đồ Use Case và Activity Diagram.", "03/08/2026 – 23/08/2026", "Phân tích yêu cầu hệ thống, xây dựng user flow. Thiết kế Use Case, Activity Diagram. Lựa chọn stack công nghệ."),
        ("2", "Thiết kế kiến trúc hệ thống Client - Server - AI Service; Thiết kế CSDL MongoDB và cấu trúc Vector Collection trong ChromaDB; Thiết kế đặc tả RESTful API.", "24/08/2026 – 13/09/2026", "Xây dựng sơ đồ kiến trúc, thiết kế database schema cho MongoDB. Chuẩn bị tập dữ liệu tri thức FAQ."),
        ("3", "Xây dựng Core Backend (FastAPI, JWT Auth, API Danh mục, Sản phẩm, Đơn hàng) và hoàn thành đánh giá Mốc 1 (M1).", "14/09/2026 – 04/10/2026", "Xây dựng RESTful API, kiểm thử Postman. Đánh giá tiến độ Mốc 1 (M1) với GVHD."),
        ("4", "Phát triển giao diện Frontend (Vue 3, Pinia): Xây dựng Trang chủ (Banner nổi bật tự chạy), Danh mục sản phẩm (Bộ lọc độc lập sticky), Chi tiết sản phẩm, Giỏ hàng, Đặt hàng và Admin Dashboard.", "05/10/2026 – 25/10/2026", "Hoàn thiện các trang chức năng client và trang quản trị admin. Tối ưu UX/UI và responsive."),
        ("5", "Xây dựng dịch vụ AI Core & Chatbot Stylist: Tích hợp RAG Pipeline, ChromaDB Vector Search, kết nối LLM (Gemini/OpenAI), kiểm thử kịch bản tư vấn outfit và size.", "26/10/2026 – 08/11/2026", "Tích hợp Chat Widget vào Frontend, xử lý prompt RAG, kiểm thử độ chính xác của phản hồi AI."),
        ("6", "Đóng gói hệ thống với Docker Compose, kiểm thử toàn diện (Function Test, Security, Performance), sửa lỗi và tối ưu hóa hệ thống.", "09/11/2026 – 22/11/2026", "Kiểm thử toàn diện luồng mua hàng, quản trị và chatbot. Tối ưu tốc độ tải trang và đóng gói Docker."),
        ("7", "Hoàn thiện toàn bộ báo cáo ĐATN, video demo, slide thuyết trình và chuẩn bị bảo vệ tốt nghiệp.", "23/11/2026 – 29/11/2026", "Hoàn thiện báo cáo đồ án, slide và sản phẩm demo để nộp hội đồng và bảo vệ.")
    ]

    for idx, (stt, kh, tg, gc) in enumerate(milestones):
        row_cells = t2.rows[idx + 1].cells
        row_cells[0].text = stt
        row_cells[1].text = kh
        row_cells[2].text = tg
        row_cells[3].text = gc

    for row in t2.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(13)

    # Update Table 3: Chữ ký Mẫu ĐATN-01
    t3 = doc.tables[3]
    t3.rows[0].cells[0].text = "Sinh viên\n(Ký, ghi rõ họ tên)"
    t3.rows[0].cells[1].text = "GV hướng dẫn\n(Ký, ghi rõ họ tên)"
    t3.rows[0].cells[2].text = "Chủ tịch hội đồng\n(Ký, ghi rõ họ tên)"
    t3.rows[1].cells[0].text = "\n\n\n"
    t3.rows[1].cells[1].text = "\n\n\n"
    t3.rows[1].cells[2].text = "\n\n\n"
    t3.rows[2].cells[0].text = "Phạm Ngọc Đức"
    t3.rows[2].cells[1].text = "Phan Hữu Trung"
    t3.rows[2].cells[2].text = ""

    for r_idx, row in enumerate(t3.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(13)
                    if r_idx in (0, 2):
                        r.bold = True

    # Update Table 5: Thông tin Mẫu ĐATN-02 (Nhật ký thực hiện)
    t5 = doc.tables[5]
    t5.rows[0].cells[1].text = "Phạm Ngọc Đức"
    t5.rows[0].cells[3].text = "0288767"
    t5.rows[1].cells[1].text = topic
    t5.rows[1].cells[3].text = "Đợt 1"
    t5.rows[2].cells[1].text = "Phan Hữu Trung"
    t5.rows[2].cells[3].text = "Công nghệ phần mềm"

    for row in t5.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(13)

    # Update Table 6: Nhật ký theo tuần (17 tuần từ 03/08/2026 đến 29/11/2026)
    t6 = doc.tables[6]
    
    # 17 weeks of diary
    weeks_data = [
        # Tuần 1
        ("Tuần 1: 03/08/2026 – 09/08/2026", "10 giờ", 
         "Nghiên cứu đề tài, tìm hiểu mô hình thương mại điện tử thời trang và giải pháp tích hợp AI Chatbot/RAG.",
         "Xác định rõ phạm vi đề tài, mục tiêu nghiên cứu và lập kế hoạch thực hiện đồ án.",
         "Khảo sát công nghệ, phân tích yêu cầu",
         "Cần tìm hiểu sâu hơn về kiến trúc RAG và Vector Database.",
         "Định hướng rõ ràng cho các bước triển khai tiếp theo."),

        # Tuần 2
        ("Tuần 2: 10/08/2026 – 16/08/2026", "12 giờ",
         "Phân tích yêu cầu hệ thống, xác định các tác nhân và xây dựng biểu đồ Use Case, Activity Diagram.",
         "Hoàn thành tài liệu đặc tả yêu cầu, sơ đồ Use Case cho Khách hàng, Admin và Chatbot AI.",
         "UML, phân tích thiết kế hệ thống",
         "Cần làm rõ hơn các luồng tương tác giữa người dùng và AI Stylist.",
         "Yêu cầu chức năng được xác định chi tiết, rõ ràng."),

        # Tuần 3
        ("Tuần 3: 17/08/2026 – 23/08/2026", "14 giờ",
         "Lựa chọn stack công nghệ (Vue 3, FastAPI, MongoDB, ChromaDB, Docker) và thiết kế kiến trúc hệ thống.",
         "Xác định mô hình kiến trúc Client - Server - AI Service, thiết lập cấu trúc mã nguồn dự án.",
         "Kiến trúc phần mềm, Docker, Web stack",
         "Cần tối ưu mô hình giao tiếp giữa Backend và AI Service.",
         "Stack công nghệ hiện đại, đáp ứng tốt mục tiêu đề tài."),

        # Tuần 4
        ("Tuần 4: 24/08/2026 – 30/08/2026", "15 giờ",
         "Thiết kế cơ sở dữ liệu MongoDB (users, products, categories, orders, carts, promotions) và ChromaDB.",
         "Hoàn thành thiết kế schema MongoDB và cấu trúc Vector Collection cho tri thức Chatbot.",
         "Thiết kế CSDL NoSQL, Vector Embeddings",
         "Cần tối ưu chỉ mục (Index) và cấu trúc thuộc tính đa dạng của sản phẩm thời trang.",
         "Cơ sở dữ liệu linh hoạt, dễ dàng mở rộng."),

        # Tuần 5
        ("Tuần 5: 31/08/2026 – 06/09/2026", "8 giờ",
         "Nghỉ lễ Quốc khánh 2/9 & Chuẩn bị môi trường phát triển Backend FastAPI, cấu hình Docker.",
         "Khởi tạo project Backend, thiết lập Dockerfile và kiểm tra kết nối MongoDB.",
         "FastAPI, Docker, Python 3.11",
         "Cần bổ sung cấu hình biến môi trường .env an toàn.",
         "Môi trường phát triển được chuẩn bị sẵn sàng."),

        # Tuần 6
        ("Tuần 6: 07/09/2026 – 13/09/2026", "16 giờ",
         "Xây dựng API Backend cơ bản: Xác thực Auth (JWT, Bcrypt), CRUD Danh mục và CRUD Sản phẩm.",
         "Hoàn thành các API Đăng nhập, Đăng ký, Quản lý sản phẩm và lọc sản phẩm.",
         "FastAPI, PyJWT, Pydantic, Motor",
         "Cần bổ sung validation chặt chẽ cho dữ liệu đầu vào.",
         "API hoạt động ổn định, kiểm thử qua Postman đạt kết quả tốt."),

        # Tuần 7
        ("Tuần 7: 14/09/2026 – 20/09/2026", "16 giờ",
         "Xây dựng API Giỏ hàng, Đơn hàng, Khuyến mãi (Voucher) và Quản lý thông tin cửa hàng.",
         "Hoàn thành toàn bộ API phục vụ quy trình mua hàng và quản trị đơn hàng.",
         "RESTful API, Business Logic",
         "Cần xử lý các trường hợp ngoại lệ khi áp dụng voucher giảm giá.",
         "Hoàn thành các nghiệp vụ cốt lõi phía Backend."),

        # Tuần 8 (Mốc 1)
        ("Tuần 8: 21/09/2026 – 27/09/2026", "15 giờ",
         "Tổng kết Mốc 1 (M1): Hoàn thiện tài liệu đặc tả, kiểm thử toàn bộ API Backend và báo cáo tiến độ với GVHD.",
         "Hoàn thành báo cáo Mốc 1, hệ thống Backend sẵn sàng tích hợp với Frontend và AI Service.",
         "Viết tài liệu kỹ thuật, Postman Test",
         "Cần tiếp thu ý kiến góp ý của GVHD để triển khai giai đoạn Frontend và AI.",
         "Hoàn thành đúng tiến độ Mốc 1 đã đề ra."),

        # Tuần 9
        ("Tuần 9: 28/09/2026 – 04/10/2026", "18 giờ",
         "Khởi tạo dự án Frontend Vue 3, cấu hình Pinia, Vue Router và xây dựng Layout chung.",
         "Hoàn thành khung ứng dụng Frontend SPA, Header cố định, Footer và trang chủ cơ bản.",
         "Vue 3, Vite, Pinia, CSS",
         "Cần tối ưu cấu trúc component để tái sử dụng hiệu quả.",
         "Giao diện chuẩn phong cách thời trang hiện đại."),

        # Tuần 10
        ("Tuần 10: 05/10/2026 – 11/10/2026", "18 giờ",
         "Phát triển trang Danh mục sản phẩm với bộ lọc Sticky cuộn độc lập và Card sản phẩm chuẩn e-commerce.",
         "Hoàn thành bộ lọc đa tiêu chí (giá, màu sắc, size, danh mục) và nút thêm giỏ hàng nhanh.",
         "Vue Component, Event Handling, CSS Flex/Grid",
         "Cần tối ưu hiệu năng render khi danh sách sản phẩm lớn.",
         "Bộ lọc hoạt động mượt mà, trải nghiệm người dùng tốt."),

        # Tuần 11
        ("Tuần 11: 12/10/2026 – 18/10/2026", "18 giờ",
         "Phát triển trang Chi tiết sản phẩm, Giỏ hàng, Thanh toán và Quản lý lịch sử đơn hàng.",
         "Hoàn thiện toàn bộ luồng mua sắm của khách hàng từ chọn hàng đến đặt hàng thành công.",
         "State Management, LocalStorage, Axios",
         "Cần kiểm tra kỹ tính toán tổng tiền và giảm giá voucher.",
         "Quy trình mua hàng hoàn chỉnh và thông suốt."),

        # Tuần 12
        ("Tuần 12: 19/10/2026 – 25/10/2026", "18 giờ",
         "Xây dựng phân hệ Admin Dashboard: Thống kê tổng quan, quản lý sản phẩm, đơn hàng, danh mục, voucher.",
         "Hoàn thiện toàn bộ giao diện và chức năng CRUD quản trị cho chủ cửa hàng.",
         "Admin UI, Data Table, Role-based Auth",
         "Cần tối ưu trải nghiệm form nhập liệu thuộc tính sản phẩm.",
         "Phân hệ quản trị trực quan, dễ thao tác."),

        # Tuần 13
        ("Tuần 13: 26/10/2026 – 01/11/2026", "18 giờ",
         "Xây dựng dịch vụ RAG AI: Xử lý nhúng vector dữ liệu sản phẩm, FAQ, chính sách vào ChromaDB.",
         "Hoàn thành pipeline nạp tri thức và truy vấn tương đồng ngữ nghĩa (Semantic Search).",
         "RAG Pipeline, ChromaDB, Embeddings",
         "Cần tối ưu việc phân đoạn văn bản (chunking) cho tài liệu chính sách.",
         "Truy xuất ngữ cảnh chính xác cao."),

        # Tuần 14
        ("Tuần 14: 02/11/2026 – 08/11/2026", "18 giờ",
         "Tích hợp mô hình LLM (Gemini/OpenAI) và hoàn thiện Chat Widget AI Stylist trên Frontend.",
         "Hoàn thành Chatbot AI có khả năng tư vấn phối đồ, chọn size và giải đáp thông tin tự nhiên.",
         "Prompt Engineering, LLM Integration, UI Chat",
         "Cần tinh chỉnh prompt để câu trả lời ngắn gọn, thân thiện và chính xác.",
         "Chatbot phản hồi nhanh, tư vấn hữu ích."),

        # Tuần 15
        ("Tuần 15: 09/11/2026 – 15/11/2026", "18 giờ",
         "Đóng gói toàn bộ hệ thống bằng Docker Compose, thiết lập môi trường chạy thử và kiểm thử liên thông.",
         "Triển khai đồng bộ Frontend, Backend, MongoDB, ChromaDB qua Docker chỉ với 1 lệnh.",
         "Docker Compose, Network, DevOps",
         "Cần tối ưu dung lượng Docker Image của Frontend và Backend.",
         "Hệ thống vận hành đồng bộ và ổn định."),

        # Tuần 16
        ("Tuần 16: 16/11/2026 – 22/11/2026", "16 giờ",
         "Kiểm thử toàn diện hệ thống, sửa các lỗi phát sinh, tối ưu hóa tốc độ tải trang và bảo mật.",
         "Hệ thống hoạt động trơn tru, không còn lỗi logic, tốc độ phản hồi nhanh.",
         "Testing, Security, Performance Tuning",
         "Cần rà soát kỹ các trường hợp biên của giỏ hàng và chatbot.",
         "Sản phẩm đạt độ hoàn thiện cao."),

        # Tuần 17
        ("Tuần 17: 23/11/2026 – 29/11/2026", "16 giờ",
         "Hoàn thiện toàn bộ tài liệu báo cáo đồ án, quay video demo, thiết kế slide thuyết trình và chuẩn bị bảo vệ.",
         "Hoàn thành quyển báo cáo ĐATN, video demo và slide thuyết trình sẵn sàng cho buổi bảo vệ.",
         "Viết báo cáo kỹ thuật, trình bày",
         "Cần luyện tập thuyết trình để tự tin trình bày trước hội đồng.",
         "Hoàn thành toàn bộ đồ án tốt nghiệp đúng hạn.")
    ]

    # Clear existing rows in Table 6 except header
    # Let's rebuild Table 6 rows cleanly
    # Header is row 0
    # For each week, we add 1 row for week title, 1 row for Sinh viên, 1 row for GVHD
    # Let's see: Table 6 currently has 29 rows. Let's adjust rows
    while len(t6.rows) > 1:
        tr = t6.rows[-1]._tr
        t6._tbl.remove(tr)

    for week_title, hours, task, result, skill, issue, review in weeks_data:
        # Row 1: Week header
        r_week = t6.add_row()
        for c in r_week.cells:
            c.text = week_title
        # Merge all cells in week header row
        a = r_week.cells[0]
        for b in r_week.cells[1:]:
            a.merge(b)
        a.text = week_title

        # Row 2: Sinh viên
        r_sv = t6.add_row()
        r_sv.cells[0].text = "Làm việc độc lập"
        r_sv.cells[1].text = hours
        r_sv.cells[2].text = task
        r_sv.cells[3].text = result
        r_sv.cells[4].text = skill
        r_sv.cells[5].text = issue
        r_sv.cells[6].text = review
        r_sv.cells[7].text = "Sinh viên"

        # Row 3: GVHD
        r_gv = t6.add_row()
        r_gv.cells[0].text = "Làm việc độc lập"
        r_gv.cells[1].text = hours
        r_gv.cells[2].text = task
        r_gv.cells[3].text = result
        r_gv.cells[4].text = skill
        r_gv.cells[5].text = issue
        r_gv.cells[6].text = review
        r_gv.cells[7].text = "GVHD"

    # Set styles for Table 6
    for row_idx, row in enumerate(t6.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(11 if row_idx > 0 else 12)
                    if row_idx == 0 or cell.text.startswith("Tuần"):
                        r.bold = True

    # Update Table 8: Mẫu ĐATN-03 (Phiếu đánh giá Mốc 1)
    t8 = doc.tables[8]
    t8.rows[0].cells[1].text = "Phạm Ngọc Đức"
    t8.rows[0].cells[3].text = "0288767"
    t8.rows[1].cells[1].text = topic
    t8.rows[1].cells[3].text = "1"
    t8.rows[2].cells[1].text = "Phan Hữu Trung"
    t8.rows[2].cells[3].text = "Công nghệ phần mềm"
    t8.rows[3].cells[1].text = "27/09/2026"
    t8.rows[3].cells[3].text = "Trực tiếp"

    for row in t8.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(13)

    # Update Table 9: Bảng đánh giá CLO
    t9 = doc.tables[9]
    t9.rows[2].cells[1].text = "Khảo sát, phân tích yêu cầu nghiệp vụ và đặc tả kiến trúc hệ thống E-commerce & AI Chatbot"
    t9.rows[2].cells[5].text = "☑"
    t9.rows[2].cells[7].text = "Đã hoàn thành phân tích yêu cầu, Use Case và Activity Diagram đầy đủ, rõ ràng."

    t9.rows[3].cells[1].text = "Thiết kế cơ sở dữ liệu MongoDB, Vector Database ChromaDB và đặc tả RESTful API"
    t9.rows[3].cells[5].text = "☑"
    t9.rows[3].cells[7].text = "Thiết kế database schema và cấu trúc tri thức RAG hợp lý, chuẩn hóa."

    t9.rows[4].cells[1].text = "Hiện thực hóa Core Backend FastAPI, xác thực JWT và kiểm thử API hoàn chỉnh"
    t9.rows[4].cells[5].text = "☑"
    t9.rows[4].cells[7].text = "Đã xây dựng và kiểm thử thành công các API cốt lõi qua Postman."

    for row in t9.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(12)

    # Update Paragraphs 96-101 for Mẫu 03
    for p in doc.paragraphs:
        if "Kết luận mốc:" in p.text:
            p.text = "Kết luận mốc: ☑ Đạt\t☐ Chưa đạt\t☐ Đạt nhưng yêu cầu chỉnh sửa"
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(13)
        if "Điểm mốc" in p.text:
            p.text = "Điểm mốc (tham khảo, nếu áp dụng):\t9.0 / 10"
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(13)
            p.runs[0].font.bold = True
        if "II. Ghi chú/yêu cầu chỉnh sửa bắt buộc" in p.text:
            pass

    # Update Table 10: Chữ ký Mẫu 03
    t10 = doc.tables[10]
    if len(t10.rows) == 2:
        t10.add_row()
    t10.rows[0].cells[0].text = "Sinh viên\n(Ký, ghi rõ họ tên)"
    t10.rows[0].cells[1].text = "GV hướng dẫn\n(Ký, ghi rõ họ tên)"
    t10.rows[0].cells[2].text = "Thành viên HĐ (nếu có)\n(Ký, ghi rõ họ tên)"
    t10.rows[1].cells[0].text = "\n\n\n"
    t10.rows[1].cells[1].text = "\n\n\n"
    t10.rows[1].cells[2].text = "\n\n\n"
    t10.rows[2].cells[0].text = "Phạm Ngọc Đức"
    t10.rows[2].cells[1].text = "Phan Hữu Trung"
    t10.rows[2].cells[2].text = ""

    for r_idx, row in enumerate(t10.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(13)
                    if r_idx in (0, 2):
                        r.bold = True

    # Save to target file
    output_path = 'e:/doan/fashion-store-ai/BaocaoM1-D1-PhamNgocDuc-0288767-67PM3.docx'
    doc.save(output_path)
    print(f"Report generated successfully at: {output_path}")

if __name__ == '__main__':
    create_report()
